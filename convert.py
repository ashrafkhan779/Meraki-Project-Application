#!/usr/bin/env python3
"""
JobOps — Job Management & Operations System (Supabase edition)
convert.py : helper utility

Commands
  python3 convert.py serve [--port 8000]      Serve index.html locally (or host it on any static web host)
  python3 convert.py test                     Unit tests for the financial formulas (mirror of the in-app engine)
  python3 convert.py template [--out jobops_import_template.xlsx]
                                              Blank Excel import template (one sheet per table) — needs: pip install openpyxl
  python3 convert.py validate <file.xlsx>     Check an Excel file offline before importing it in the app

The column specification (TABLES) below is the single source of truth: build.py embeds it in index.html,
so the app's Excel export/import and this template always share the same layout.
"""
import argparse
import datetime as dt
import http.server
import json
import os
import socketserver
import sys
from decimal import Decimal, ROUND_HALF_UP

HERE = os.path.dirname(os.path.abspath(__file__))

# ---------------------------------------------------------------------------
# Table specification — single source for template, import, validation, SQL.
# type: text | num | int | date | ts | bool | arr | json ; fk: referenced table
# ---------------------------------------------------------------------------
C = lambda name, typ="text", fk=None, req=False: {"name": name, "type": typ, "fk": fk, "req": req}

TABLES = {
    "customers": [C("id", req=True), C("code", req=True), C("name", req=True), C("contactPerson"), C("email"), C("phone"),
                  C("address"), C("trn"), C("paymentTerms"), C("creditLimit", "num"), C("zohoCustomerId"), C("notes"), C("status")],
    "endClients": [C("id", req=True), C("code", req=True), C("name", req=True), C("contactPerson"), C("email"), C("phone"),
                   C("address"), C("customerIds", "arr"), C("notes"), C("status")],
    "suppliers": [C("id", req=True), C("code", req=True), C("name", req=True), C("contactPerson"), C("email"), C("phone"),
                  C("address"), C("trn"), C("paymentTerms"), C("category"), C("notes"), C("status")],
    "employees": [C("id", req=True), C("code", req=True), C("name", req=True), C("type", req=True), C("phone"), C("email"),
                  C("title"), C("department"), C("joiningDate", "date"), C("status"), C("notes"), C("employmentType"),
                  C("stdHoursPerDay", "num"), C("costRate", "num"), C("skills"), C("supplierId", fk="suppliers"),
                  C("rateBasis"), C("agreedRate", "num"), C("capacityPct", "num"), C("contractStart", "date"), C("contractEnd", "date")],
    "leaves": [C("id", req=True), C("employeeId", fk="employees", req=True), C("fromDate", "date", req=True),
               C("toDate", "date", req=True), C("type"), C("notes")],
    "inventoryItems": [C("id", req=True), C("sku", req=True), C("name", req=True), C("uom"), C("category"), C("store"),
                       C("qtyOnHand", "num"), C("avgCost", "num"), C("status")],
    "stockReceipts": [C("id", req=True), C("itemId", fk="inventoryItems", req=True), C("receiptDate", "date"),
                      C("qty", "num"), C("unitCost", "num"), C("ref")],
    "jobs": [C("id", req=True), C("jobNo", req=True), C("type", req=True), C("openingDate", "date", req=True),
             C("customerId", fk="customers"), C("endClientId", fk="endClients"), C("zohoSo"), C("zohoSoId"), C("description"),
             C("status", req=True), C("completionPct", "num"), C("lpoDate", "date"), C("customerOrderNo"), C("poValue", "num"),
             C("actualPoValue", "num"), C("psDate", "date"), C("pfDate", "date"), C("actualStart", "date"),
             C("actualFinish", "date"), C("projectEngineerId", fk="employees"), C("serviceEngineerId", fk="employees"),
             C("comments"), C("locked", "bool"), C("closureDate", "date"), C("closureComments"), C("closedBy", fk="users"),
             C("closedAt", "ts")],
    "plannedManpower": [C("id", req=True), C("jobId", fk="jobs", req=True), C("employeeId", fk="employees"), C("resourceName"),
                        C("resourceType"), C("supplierId", fk="suppliers"), C("headcount", "num"), C("hoursPerPerson", "num"),
                        C("days", "num"), C("rate", "num"), C("startDate", "date"), C("endDate", "date"), C("comments"),
                        C("assignmentId", fk="assignments")],
    "actualManpower": [C("id", req=True), C("jobId", fk="jobs", req=True), C("employeeId", fk="employees"), C("resourceName"),
                       C("resourceType"), C("supplierId", fk="suppliers"), C("headcount", "num"), C("hoursPerPerson", "num"),
                       C("days", "num"), C("rate", "num"), C("dateFrom", "date"), C("dateTo", "date"), C("comments"),
                       C("assignmentId", fk="assignments")],
    "assignments": [C("id", req=True), C("employeeId", fk="employees", req=True), C("jobId", fk="jobs", req=True),
                    C("startDate", "date", req=True), C("endDate", "date", req=True), C("plannedHours", "num"),
                    C("allocationPct", "num"), C("status"), C("comments"), C("overrideBy", fk="users"), C("overrideReason")],
    "plannedMaterials": [C("id", req=True), C("jobId", fk="jobs", req=True), C("description"), C("category"),
                         C("supplierId", fk="suppliers"), C("qty", "num"), C("uom"), C("unitPrice", "num"),
                         C("expectedDate", "date"), C("comments")],
    "actualMaterials": [C("id", req=True), C("jobId", fk="jobs", req=True), C("description"), C("category"),
                        C("supplierId", fk="suppliers"), C("qty", "num"), C("uom"), C("unitPrice", "num"),
                        C("purchaseDate", "date"), C("supplierRef"), C("comments")],
    "inventoryTransfers": [C("id", req=True), C("jobId", fk="jobs", req=True), C("itemId", fk="inventoryItems", req=True),
                           C("qty", "num"), C("uom"), C("transferDate", "date"), C("issuedBy"), C("receivedBy"), C("store"),
                           C("costingMethod"), C("unitCost", "num"), C("costSource"), C("ref", req=True), C("status"),
                           C("comments")],
    "plannedMisc": [C("id", req=True), C("jobId", fk="jobs", req=True), C("category"), C("description"),
                    C("supplierId", fk="suppliers"), C("payee"), C("expenseDate", "date"), C("amount", "num"), C("comments")],
    "actualMisc": [C("id", req=True), C("jobId", fk="jobs", req=True), C("category"), C("description"),
                   C("supplierId", fk="suppliers"), C("payee"), C("expenseDate", "date"), C("amount", "num"),
                   C("paymentMode"), C("paymentRef"), C("attachmentName"), C("comments")],
    "miscPayments": [C("id", req=True), C("miscId", fk="actualMisc", req=True), C("payDate", "date"), C("amount", "num"),
                     C("mode"), C("ref")],
    "invoices": [C("id", req=True), C("jobId", fk="jobs", req=True), C("invoiceNo", req=True), C("invoiceDate", "date"),
                 C("dueDate", "date"), C("amount", "num"), C("status"), C("zohoInvoiceId"), C("notes")],
    "invoicePayments": [C("id", req=True), C("invoiceId", fk="invoices", req=True), C("payDate", "date"), C("amount", "num"),
                        C("mode"), C("ref"), C("zohoPaymentId")],
    "jobStatusHistory": [C("id", req=True), C("jobId", fk="jobs", req=True), C("fromStatus"), C("toStatus"), C("at", "ts"),
                         C("userId", fk="users"), C("note")],
    "poRevisions": [C("id", req=True), C("jobId", fk="jobs", req=True), C("field"), C("fromValue", "num"),
                    C("toValue", "num"), C("at", "ts"), C("userId", fk="users"), C("reason")],
    "jobComments": [C("id", req=True), C("jobId", fk="jobs", req=True), C("userId", fk="users"), C("at", "ts"), C("text")],
    "jobAttachments": [C("id", req=True), C("jobId", fk="jobs", req=True), C("name"), C("size", "int"), C("mime"),
                       C("at", "ts"), C("userId", fk="users"), C("note")],
    "closures": [C("id", req=True), C("jobId", fk="jobs", req=True), C("closureDate", "date"), C("comments"),
                 C("finalStatus"), C("financialsConfirmed", "bool"), C("pendingConfirmed", "bool"),
                 C("closedBy", fk="users"), C("closedAt", "ts"), C("reopenedBy", fk="users"), C("reopenedAt", "ts"),
                 C("reopenReason")],
    "roles": [C("id", req=True), C("name", req=True), C("system", "bool"), C("jobTypes", "arr"), C("permissions", "json")],
    "users": [C("id", req=True), C("name", req=True), C("email", req=True), C("roleId", fk="roles", req=True),
              C("department"), C("active", "bool")],
    "auditLog": [C("id", req=True), C("at", "ts"), C("userId", fk="users"), C("action"), C("entity"), C("recordId"),
                 C("before", "json"), C("after", "json"), C("reason")],
}
# Tables that carry createdAt/createdBy/updatedAt/version/archived
IMPORTABLE = [t for t in TABLES if t != "auditLog"]

PERMISSIONS = [
    "viewOverview", "viewJobs", "viewRoster", "viewMaster", "viewReports", "viewSettings",
    "jobCreate", "quickAddMaster", "jobEdit", "jobDelete", "viewFinancials", "editFinancials", "viewCostRates",
    "manageManpower", "manageMaterials", "manageMisc", "changeInvoiceStatus", "recordPayments",
    "closeJobs", "reopenJobs", "editClosedJobs", "exportReports", "manageMaster", "rosterAssign",
    "overrideConflicts", "manageUsers", "manageSettings",
]
JOB_TYPES = ["Service", "Supply", "Supply/Services"]


def role(rid, name, perms, types=None):
    p = {k: (k in perms) for k in PERMISSIONS}
    return {"id": rid, "name": name, "system": True, "jobTypes": types or list(JOB_TYPES), "permissions": p}


ALL = set(PERMISSIONS)
VIEW_TABS = {"viewOverview", "viewJobs", "viewRoster", "viewMaster", "viewReports"}
DEFAULT_ROLES = [
    role("role_super", "Super Admin", ALL),
    role("role_admin", "Administrator", ALL - {"manageUsers"} | {"manageUsers"}),
    role("role_mgmt", "Management", VIEW_TABS | {"viewFinancials", "viewCostRates", "exportReports"}),
    role("role_fin", "Finance / Accounts", VIEW_TABS | {"viewFinancials", "editFinancials", "viewCostRates", "manageMisc",
                                                       "changeInvoiceStatus", "recordPayments", "closeJobs",
                                                       "exportReports"}),
    role("role_pm", "Project Manager", VIEW_TABS | {"jobCreate", "quickAddMaster", "jobEdit", "viewFinancials", "editFinancials",
                                                    "manageManpower", "manageMaterials", "manageMisc", "rosterAssign",
                                                    "exportReports"}),
    role("role_ops", "Operations", VIEW_TABS - {"viewReports"} | {"jobEdit", "manageMaterials", "manageMisc"}),
    role("role_hr", "HR / Resource Planner", {"viewOverview", "viewJobs", "viewRoster", "viewMaster", "rosterAssign",
                                              "manageManpower", "manageMaster", "overrideConflicts"}),
    role("role_ro", "Read-Only User", {"viewOverview", "viewJobs", "viewRoster", "viewReports"}),
]

DEFAULT_SETTINGS = {
    "companyName": "Your Company LLC",
    "logoText": "JO",
    "currency": "AED",
    "precision": 2,
    "locale": "en-AE",
    "jobNumber": {"prefix": "JOB", "pattern": "{PREFIX}-{TYPE}-{YYYY}-{SEQ}", "padding": 4, "nextSeq": 1,
                  "typeCodes": {"Service": "SRV", "Supply": "SUP", "Supply/Services": "SSV"}},
    "jobStatuses": ["Open", "Planned", "In Progress", "On Hold", "Completed", "Closed", "Cancelled"],
    "expenseCategories": ["Transportation", "Accommodation", "Permits & Gate Passes", "Equipment Rental", "Delivery",
                          "Parking & Tolls", "Meals", "Other"],
    "materialCategories": ["Consumables", "Spare Parts", "Electrical", "Mechanical", "Piping & Valves", "Safety / PPE",
                           "Tools", "Other"],
    "paymentModes": ["Cash", "Bank Transfer", "Company Card", "Cheque", "Other"],
    "uoms": ["EA", "PCS", "SET", "M", "KG", "L", "BOX", "LOT"],
    "departments": ["Operations", "Projects", "Workshop", "Logistics", "Finance", "HR", "Management"],
    "workingDays": [1, 2, 3, 4, 5],
    "hoursPerDay": 8,
    "timezone": "Asia/Dubai",
    "fiscalYearStartMonth": 1,
    "revenueRule": "po_value",
    "inventoryCosting": "weighted_average",
    "suggest100OnCompleted": True,
    "schemaVersion": 2,
    "gpWarnPct": 15,
    "attachmentLimitMb": 10,
    "notifyOnClose": True,
    "notifyOnOverbook": True,
    "exportRequiresPermission": True,
}

# ---------------------------------------------------------------------------
# Financial calculation policy (mirrors index.html → Calc)
# Each line total is rounded to the currency minor unit (2 dp, half-up), totals are sums of rounded lines.
# ---------------------------------------------------------------------------
TWO = Decimal("0.01")


def D(x):
    if x is None or x == "":
        return Decimal("0")
    return Decimal(str(x))


def r2(x):
    return D(x).quantize(TWO, rounding=ROUND_HALF_UP)


def line_days(l):
    d = l.get("days")
    return D(d) if d not in (None, "") and D(d) > 0 else Decimal(1)


def line_mp(l):
    """headcount × hours per person per day × days × hourly rate"""
    return r2(D(l.get("headcount")) * D(l.get("hoursPerPerson")) * line_days(l) * D(l.get("rate")))


def working_days(a, b, wd=(1, 2, 3, 4, 5)):
    """working days in [a, b] (ISO weekday numbers 1=Mon); a range with none counts calendar days"""
    if not a:
        return None
    b = b if b and b >= a else a
    cnt = cal = 0
    d = a
    while d <= b:
        cal += 1
        if d.isoweekday() in wd:
            cnt += 1
        d += dt.timedelta(days=1)
    return cnt or cal


def line_qty(l, price="unitPrice"):
    return r2(D(l.get("qty")) * D(l.get(price)))


def job_financials(db, job):
    jid = job["id"]
    by = lambda t: [x for x in db.get(t, []) if x.get("jobId") == jid and not x.get("archived")]
    pm = sum((line_mp(l) for l in by("plannedManpower")), Decimal(0))
    am = sum((line_mp(l) for l in by("actualManpower")), Decimal(0))
    pmat = sum((line_qty(l) for l in by("plannedMaterials")), Decimal(0))
    amat_p = sum((line_qty(l) for l in by("actualMaterials")), Decimal(0))
    amat_t = sum((line_qty(l, "unitCost") for l in by("inventoryTransfers") if l.get("status", "Posted") == "Posted"),
                 Decimal(0))
    pmisc = sum((r2(l.get("amount")) for l in by("plannedMisc")), Decimal(0))
    amisc = sum((r2(l.get("amount")) for l in by("actualMisc")), Decimal(0))
    rule = db.get("settings", {}).get("revenueRule", "po_value")
    ppo = r2(job.get("poValue"))
    if rule == "invoiced":
        apo = sum((r2(i.get("amount")) for i in by("invoices")), Decimal(0))
    else:
        apo = r2(job["actualPoValue"]) if job.get("actualPoValue") not in (None, "") else ppo
    ptot = pm + pmat + pmisc
    atot = am + amat_p + amat_t + amisc
    pgp, agp = ppo - ptot, apo - atot
    pct = lambda gp, po: (gp / po * 100).quantize(TWO, rounding=ROUND_HALF_UP) if po > 0 else None
    return {"plannedPO": ppo, "actualPO": apo, "plannedManpower": pm, "actualManpower": am,
            "plannedMaterial": pmat, "actualMaterialPurchased": amat_p, "actualMaterialTransfer": amat_t,
            "actualMaterial": amat_p + amat_t, "plannedMisc": pmisc, "actualMisc": amisc,
            "plannedCost": ptot, "actualCost": atot, "plannedGP": pgp, "actualGP": agp,
            "plannedGPPct": pct(pgp, ppo), "actualGPPct": pct(agp, apo), "costVariance": atot - ptot,
            "gpVariance": agp - pgp}


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------
def empty_db():
    db = {"settings": json.loads(json.dumps(DEFAULT_SETTINGS)),
          "roles": json.loads(json.dumps(DEFAULT_ROLES))}
    for t in TABLES:
        if t not in ("roles",):
            db.setdefault(t, [])
    return db


# ---------------------------------------------------------------------------
# Validation
# ---------------------------------------------------------------------------
def validate(db, verbose=True):
    errors, warns = [], []
    idx = {t: {r.get("id"): r for r in db.get(t, [])} for t in TABLES}
    for t, cols in TABLES.items():
        rows = db.get(t, [])
        seen = set()
        for r in rows:
            rid = r.get("id")
            if rid in seen:
                errors.append(f"{t}: duplicate id {rid}")
            seen.add(rid)
            for c in cols:
                v = r.get(c["name"])
                if c["req"] and (v is None or v == ""):
                    errors.append(f"{t}.{c['name']} required (id={rid})")
                if c["fk"] and v not in (None, ""):
                    if v not in idx[c["fk"]]:
                        errors.append(f"{t}.{c['name']} → {c['fk']} missing '{v}' (id={rid})")
                if c["type"] == "date" and v:
                    try:
                        dt.date.fromisoformat(v)
                    except Exception:
                        errors.append(f"{t}.{c['name']} invalid date '{v}' (id={rid})")
                if c["type"] == "num" and v not in (None, ""):
                    try:
                        if D(v) < 0:
                            errors.append(f"{t}.{c['name']} negative ({v}) (id={rid})")
                    except Exception:
                        errors.append(f"{t}.{c['name']} not numeric '{v}' (id={rid})")
    nos = [j["jobNo"] for j in db.get("jobs", [])]
    for n in set(nos):
        if nos.count(n) > 1:
            errors.append(f"jobs: duplicate jobNo {n}")
    for j in db.get("jobs", []):
        if j.get("type") not in JOB_TYPES:
            errors.append(f"jobs: invalid type '{j.get('type')}' ({j['jobNo']})")
        p = j.get("completionPct")
        if p not in (None, "") and not (0 <= float(p) <= 100):
            errors.append(f"jobs: completionPct out of range ({j['jobNo']})")
        for a, b in (("psDate", "pfDate"), ("actualStart", "actualFinish")):
            if j.get(a) and j.get(b) and j[b] < j[a]:
                errors.append(f"jobs: {b} before {a} ({j['jobNo']})")
        if j.get("status") == "Closed" and not j.get("locked"):
            warns.append(f"jobs: {j['jobNo']} is Closed but not locked")
    refs = [t["ref"] for t in db.get("inventoryTransfers", []) if t.get("status") == "Posted"]
    for r in set(refs):
        if refs.count(r) > 1:
            errors.append(f"inventoryTransfers: transfer ref {r} posted twice (double charge)")
    for inv in db.get("invoices", []):
        paid = sum((r2(p["amount"]) for p in db.get("invoicePayments", []) if p["invoiceId"] == inv["id"]), Decimal(0))
        exp = "Paid" if paid >= r2(inv["amount"]) and paid > 0 else ("Partially Paid" if paid > 0 else "Unpaid")
        if inv.get("status") != exp:
            warns.append(f"invoices: {inv['invoiceNo']} status '{inv.get('status')}' but payments imply '{exp}'")
    if verbose:
        print(f"Tables: " + ", ".join(f"{t}={len(db.get(t, []))}" for t in TABLES if db.get(t)))
        for e in errors:
            print("  ERROR ", e)
        for w in warns:
            print("  WARN  ", w)
        print(f"{len(errors)} error(s), {len(warns)} warning(s)")
    return errors, warns


def print_financials(db):
    cur = db["settings"].get("currency", "AED")
    print(f"\n{'Job':<22}{'Status':<13}{'PO':>13}{'Plan cost':>13}{'Act cost':>13}{'Plan GP%':>9}{'Act GP%':>9}")
    for j in db["jobs"]:
        f = job_financials(db, j)
        fp = lambda x: "N/A" if x is None else f"{x:.1f}"
        print(f"{j['jobNo']:<22}{j['status']:<13}{f['plannedPO']:>13,.2f}{f['plannedCost']:>13,.2f}"
              f"{f['actualCost']:>13,.2f}{fp(f['plannedGPPct']):>9}{fp(f['actualGPPct']):>9}")
    print(f"(amounts in {cur})")


# ---------------------------------------------------------------------------
# Tests
# ---------------------------------------------------------------------------
def run_tests():
    ok = 0
    fails = []

    def check(name, cond):
        nonlocal ok
        if cond:
            ok += 1
            print(f"  PASS  {name}")
        else:
            fails.append(name)
            print(f"  FAIL  {name}")

    db = empty_db()
    job = {"id": "j1", "poValue": 10000, "actualPoValue": None}
    db["jobs"] = [job]
    db["plannedMisc"] = [{"id": "a", "jobId": "j1", "amount": 6000}]
    f = job_financials(db, job)
    check("T4 planned GP = 4,000", f["plannedGP"] == Decimal("4000.00"))
    check("T4 planned GP% = 40%", f["plannedGPPct"] == Decimal("40.00"))
    db["actualMisc"] = [{"id": "b", "jobId": "j1", "amount": 7000}]
    f = job_financials(db, job)
    check("T5 actual GP = 3,000", f["actualGP"] == Decimal("3000.00"))
    check("T5 actual GP% = 30%", f["actualGPPct"] == Decimal("30.00"))
    check("T5 cost variance = +1,000 (unfavourable)", f["costVariance"] == Decimal("1000.00"))
    job["poValue"] = 0
    f = job_financials(db, job)
    check("T13 zero PO → GP% N/A (no division error)", f["plannedGPPct"] is None and f["actualGPPct"] is None)
    job["poValue"] = None
    check("T13 blank PO handled", job_financials(db, job)["plannedPO"] == Decimal("0.00"))
    check("Manpower line = headcount × hours/person × rate (no days → 1)", line_mp({"headcount": 3, "hoursPerPerson": 8, "rate": 45.5}) == Decimal("1092.00"))
    check("Manpower line × days: 2 × 8 h × 5 days × 50 = 4,000", line_mp({"headcount": 2, "hoursPerPerson": 8, "days": 5, "rate": 50}) == Decimal("4000.00"))
    check("Working days Mon 12 Oct – Sun 18 Oct 2026 = 5", working_days(dt.date(2026, 10, 12), dt.date(2026, 10, 18)) == 5)
    check("Decimal amounts round half-up per line", line_qty({"qty": 3, "unitPrice": 0.335}) == Decimal("1.01"))
    db2 = empty_db()
    j2 = {"id": "j2", "poValue": 1000}
    db2["jobs"] = [j2]
    db2["inventoryTransfers"] = [{"id": "t1", "jobId": "j2", "qty": 10, "unitCost": 5, "status": "Posted", "ref": "A"},
                                 {"id": "t2", "jobId": "j2", "qty": 10, "unitCost": 5, "status": "Reversed", "ref": "B"}]
    db2["actualMaterials"] = [{"id": "m", "jobId": "j2", "qty": 2, "unitPrice": 100}]
    f2 = job_financials(db2, j2)
    check("Internal transfer costed (not free) and reversed transfer excluded",
          f2["actualMaterialTransfer"] == Decimal("50.00") and f2["actualMaterial"] == Decimal("250.00"))
    j2["actualPoValue"] = 1200
    check("Revised actual PO value used for actual GP", job_financials(db2, j2)["actualPO"] == Decimal("1200.00"))
    db2["settings"]["revenueRule"] = "invoiced"
    db2["invoices"] = [{"id": "i", "jobId": "j2", "amount": 900}]
    check("Revenue rule 'invoiced' uses invoice total", job_financials(db2, j2)["actualPO"] == Decimal("900.00"))
    print(f"\n{ok} passed, {len(fails)} failed")
    return not fails


# ---------------------------------------------------------------------------
# Excel template / import
# ---------------------------------------------------------------------------
def need_openpyxl():
    try:
        import openpyxl  # noqa
        return openpyxl
    except ImportError:
        sys.exit("openpyxl is required for this command:  pip install openpyxl")


def make_template(out):
    ox = need_openpyxl()
    from openpyxl.styles import Font, PatternFill
    wb = ox.Workbook()
    ws = wb.active
    ws.title = "README"
    lines = ["JobOps import template", "",
             "One sheet per table. Row 1 = column keys (do not rename). Leave unknown cells blank.",
             "id: any unique text (e.g. cus_001). Foreign keys (…Id columns) must reference an id on the related sheet.",
             "Dates: YYYY-MM-DD.  Numbers: plain numbers, no currency symbols.  customerIds: comma-separated ids.",
             "jobs.type: Service | Supply | Supply/Services.   employees.type: inhouse | outsourced.",
             "Import it in the app: Settings → Data & integrations → Upload Excel.  Check offline first: python3 convert.py validate this_file.xlsx"]
    for i, l in enumerate(lines, 1):
        ws.cell(row=i, column=1, value=l)
    ws["A1"].font = Font(bold=True, size=14)
    for t in IMPORTABLE:
        sh = wb.create_sheet(t[:31])
        for ci, c in enumerate(TABLES[t], 1):
            cell = sh.cell(row=1, column=ci, value=c["name"])
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="244DC7" if c["req"] else "4F607E")
            sh.column_dimensions[cell.column_letter].width = max(12, len(c["name"]) + 4)
        sh.freeze_panes = "A2"
    wb.save(out)
    print(f"wrote {out}  (dark-blue headers = required)")


def coerce(v, typ):
    if v is None or (isinstance(v, str) and v.strip() == ""):
        return [] if typ == "arr" else None
    if typ == "num":
        return float(v)
    if typ == "int":
        return int(float(v))
    if typ == "bool":
        return str(v).strip().lower() in ("1", "true", "yes", "y")
    if typ == "date":
        if isinstance(v, (dt.datetime, dt.date)):
            return (v.date() if isinstance(v, dt.datetime) else v).isoformat()
        return dt.date.fromisoformat(str(v).strip()[:10]).isoformat()
    if typ == "arr":
        return [s.strip() for s in str(v).split(",") if s.strip()]
    return str(v).strip()


def read_excel(path):
    """Excel workbook (app export / template layout) -> dataset dict"""
    ox = need_openpyxl()
    db = empty_db()
    wb = ox.load_workbook(path, data_only=True)
    for t in TABLES:
        if t not in wb.sheetnames:
            continue
        rows = list(wb[t].iter_rows(values_only=True))
        if len(rows) < 2:
            continue
        hdr = [str(h).strip() if h else "" for h in rows[0]]
        types = {c["name"]: c["type"] for c in TABLES[t]}
        for row in rows[1:]:
            if not any(v not in (None, "") for v in row):
                continue
            rec = {}
            for h, v in zip(hdr, row):
                if h:
                    rec[h] = coerce(v, types.get(h, "text"))
            if rec.get("id"):
                db[t].append(rec)
    if "Settings" in wb.sheetnames:
        for key, val, typ in [r[:3] for r in wb["Settings"].iter_rows(values_only=True, min_row=2)]:
            if key:
                db["settings"][key] = json.loads(val) if typ == "json" and isinstance(val, str) else val
    return db


def validate_excel(path):
    db = read_excel(path)
    errors, _ = validate(db)
    if db.get("jobs"):
        print_financials(db)
    sys.exit(1 if errors else 0)


# ---------------------------------------------------------------------------
# Server
# ---------------------------------------------------------------------------
def serve(port):
    os.chdir(HERE)

    class H(http.server.SimpleHTTPRequestHandler):
        def end_headers(self):
            self.send_header("Cache-Control", "no-store")
            super().end_headers()

        def log_message(self, *a):
            pass

    socketserver.TCPServer.allow_reuse_address = True
    with socketserver.TCPServer(("127.0.0.1", port), H) as httpd:
        print(f"JobOps →  http://localhost:{port}/index.html   (Ctrl+C to stop)")
        try:
            httpd.serve_forever()
        except KeyboardInterrupt:
            print("\nstopped")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sp = ap.add_subparsers(dest="cmd", required=True)
    p = sp.add_parser("serve"); p.add_argument("--port", type=int, default=8000)
    sp.add_parser("test")
    p = sp.add_parser("template"); p.add_argument("--out", default="jobops_import_template.xlsx")
    p = sp.add_parser("validate"); p.add_argument("file")
    a = ap.parse_args()
    if a.cmd == "serve":
        serve(a.port)
    elif a.cmd == "test":
        sys.exit(0 if run_tests() else 1)
    elif a.cmd == "template":
        make_template(a.out)
    elif a.cmd == "validate":
        validate_excel(a.file)


if __name__ == "__main__":
    main()
